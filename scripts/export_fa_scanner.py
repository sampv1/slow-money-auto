#!/usr/bin/env python3
"""
export_fa_scanner.py — the FA Scanner's data for every symbol, to Excel.

ONE SCRIPT FOR BOTH SIDES OF A COMPARISON. It is written to be run before and
after a scoring change and diffed, so it must read nothing that the change
itself rewrites: every column comes from the stored rows, none is recomputed
here. A second implementation of the rubric in an export would make the diff a
comparison of two exporters rather than of two scorings.

Three sheets, because the FA Scanner is three rubrics and only one of them has
C1/C2/C3:

  manufacturing   `fa_scores`, minus real estate and securities — the 9-criterion
                  rubric, the only one an EPS change can touch.
  real_estate     `fa_re_scores`, the 13-criterion BĐS rubric.
  securities      `fa_securities_scores` at the active model, one row per broker
                  for the newest session.

Usage:
  python3 export_fa_scanner.py --out ../data/exports/fa_before.xlsx
  python3 export_fa_scanner.py --out ../data/exports/fa_after.xlsx --period 2026-Q2
"""

import argparse
import datetime as dt
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from ta.common import get_supabase_client, paged_select, safe_execute  # noqa: E402

# Every fa_scores column a reader of the scanner could want, in the order the
# page presents them. Spelled out rather than `*` so a schema addition cannot
# silently change the shape of a comparison already in flight.
MFG_COLS = [
    "symbol", "as_of_period",
    "c1_eps_yoy", "c1_pts",
    "c2_eps_3q_avg_yoy", "c2_pts",
    "c3_eps_pos_count", "c3_pts",
    "c4_rev_yoy", "c4_pts",
    "c5_gross_margin_delta", "c5_pts",
    "c6_net_margin_delta", "c6_pts",
    "c7_roe", "c7_pts",
    "c8_debt_to_equity", "c8_pts",
    "c9_current_pe", "c9_pts",
    "total_score", "normalized_score", "rating",
    "final_score", "final_grade",
    "current_eps_ttm", "current_pe", "pe_5y_median",
    "current_price", "current_price_date", "notes",
]


def page(client, table: str, cols: str, narrow=None, order=("symbol",)):
    def build(off, lim):
        q = client.table(table).select(cols)
        if narrow:
            q = narrow(q)
        for c in order:
            q = q.order(c)
        return q.range(off, off + lim - 1)
    return paged_select(build, label=f"{table} export")


def latest_period(client) -> str:
    data = safe_execute(client.rpc("fa_quarters"), label="fa_quarters").data or []
    qs = [q for q in data if isinstance(q, str)]
    if not qs:
        raise RuntimeError("fa_quarters returned nothing")
    return sorted(qs)[-1]


def load_manufacturing(client, period: str) -> list[dict]:
    rows = page(client, "fa_scores", ",".join(MFG_COLS),
                narrow=lambda q: q.eq("as_of_period", period))
    # The page subtracts the other two rubrics — a property developer carries a
    # stale manufacturing row, and a broker's row is scored on criteria that do
    # not describe its income statement (see the Final Score notes).
    ind = {r["symbol"]: r.get("industry_group") for r in page(
        client, "fa_industry", "symbol,industry_group")}
    prof = {r["symbol"]: r.get("com_type_code") for r in page(
        client, "symbol_profile", "symbol,com_type_code")}
    names = {r["symbol"]: r.get("short_name_vi") or r.get("name_vi") for r in page(
        client, "symbol_profile", "symbol,short_name_vi,name_vi")}
    rel = {}
    for r in page(client, "fa_statement_release_dates", "symbol,period,release_date",
                  narrow=lambda q: q.eq("period", period), order=("symbol",)):
        rel[r["symbol"]] = r.get("release_date")

    out = []
    for r in rows:
        s = r["symbol"]
        if ind.get(s) == "real_estate" or prof.get(s) == "CK":
            continue
        out.append({"symbol": s, "name": names.get(s), "release_date": rel.get(s),
                    **{k: r.get(k) for k in MFG_COLS if k != "symbol"}})
    return sorted(out, key=lambda r: r["symbol"])


def load_real_estate(client, period: str) -> list[dict]:
    """Every stored RE row, not just `period`.

    `fa_re_scores` holds only the quarters BA has exported (2026-Q2 today), so
    filtering on the manufacturing quarter would return nothing the first time
    the two diverge — and an empty sheet reads as "the rubric lost its data"
    rather than "it is scored on its own calendar".
    """
    rows = page(client, "fa_re_scores",
                "symbol,as_of_period,total_score,normalized_score,n_scored,scorable_weight")
    return sorted(rows, key=lambda r: (r.get("as_of_period") or "", r["symbol"]))


def load_securities(client) -> list[dict]:
    from fa.securities import MODEL_VERSION as ACTIVE_MODEL  # what the dashboard pins
    dates = page(client, "fa_securities_scores", "as_of_date",
                 narrow=lambda q: q.eq("model_version", ACTIVE_MODEL),
                 order=("as_of_date",))
    if not dates:
        return []
    newest = max(d["as_of_date"] for d in dates)
    rows = page(client, "fa_securities_scores",
                "symbol,as_of_date,model_version,final_fa_score,provisional_score,"
                "final_available_max,final_coverage,score_status,data_group,fa_status",
                narrow=lambda q: q.eq("model_version", ACTIVE_MODEL).eq("as_of_date", newest))
    return sorted(rows, key=lambda r: r["symbol"])


def write_xlsx(path: Path, sheets: dict[str, list[dict]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        if not rows:
            ws.append(["(no rows)"])
            continue
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for r in rows:
            ws.append([r.get(h) for h in headers])
        for i, h in enumerate(headers, start=1):
            width = max(len(str(h)), *(len(str(r.get(h) or "")) for r in rows[:400]))
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width + 2, 40)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser(description="Export the FA Scanner's data to Excel")
    ap.add_argument("--out", required=True)
    ap.add_argument("--period", help="quarter to export (default: the newest scored)")
    args = ap.parse_args()

    client = get_supabase_client()
    period = args.period or latest_period(client)
    print(f"Exporting FA Scanner data for {period}")

    mfg = load_manufacturing(client, period)
    re_rows = load_real_estate(client, period)
    sec = load_securities(client)
    print(f"  manufacturing {len(mfg):,} · real estate {len(re_rows):,} · securities {len(sec):,}")

    meta = [{"key": "exported_at", "value": dt.datetime.now().isoformat(timespec="seconds")},
            {"key": "period", "value": period},
            {"key": "manufacturing_rows", "value": len(mfg)},
            {"key": "real_estate_rows", "value": len(re_rows)},
            {"key": "securities_rows", "value": len(sec)}]

    out = Path(args.out)
    write_xlsx(out, {"manufacturing": mfg, "real_estate": re_rows,
                     "securities": sec, "meta": meta})
    print(f"Wrote {out} ({out.stat().st_size / 1024:.0f} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
